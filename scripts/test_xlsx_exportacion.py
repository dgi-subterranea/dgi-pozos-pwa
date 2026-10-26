"""Verifica desde afuera (openpyxl) el XLSX de exportacion de Mi seleccion que genera el codigo REAL del frontend
(scripts/generar_xlsx_exportacion.js -> js/exportacionSeleccionLogic.js + js/xlsxLogic.js + JSZip).

No usa LibreOffice. Se omite (skip) si faltan node u openpyxl. Correr desde scripts/:
    python -m unittest test_xlsx_exportacion
"""
import json
import os
import shutil
import subprocess
import tempfile
import unittest
import zipfile

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

AQUI = os.path.dirname(os.path.abspath(__file__))
GENERADOR = os.path.join(AQUI, 'generar_xlsx_exportacion.js')
NODE = shutil.which('node')


def generar(directorio, modo):
    destino = os.path.join(directorio, 'export_%s.xlsx' % modo)
    r = subprocess.run([NODE, GENERADOR, destino, modo], capture_output=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.decode('utf-8', 'replace'))
    return destino, json.loads(r.stdout.decode('utf-8'))


def esperado_texto(valor):
    """Lo que el escritor garantiza para un texto: CR/CRLF -> LF y sin caracteres de control ilegales."""
    if not isinstance(valor, str):
        return valor
    s = valor.replace('\r\n', '\n').replace('\r', '\n')
    return ''.join(c for c in s if c in '\t\n' or ord(c) >= 0x20)


@unittest.skipIf(openpyxl is None or NODE is None, 'faltan openpyxl o node')
class TestXlsxExportacion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.mkdtemp(prefix='xlsx_exp_')
        cls.ruta, cls.modelo = generar(cls.dir, 'muestra')
        cls.wb = openpyxl.load_workbook(cls.ruta)             # formulas SIN evaluar: si hubiera alguna, data_type seria 'f'

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, ignore_errors=True)

    def test_el_libro_abre_y_tiene_las_dos_hojas(self):
        self.assertEqual(self.wb.sheetnames, ['Pozos', 'Información'])

    def test_encabezado_de_pozos(self):
        ws = self.wb['Pozos']
        encabezados = [c.value for c in ws[1]]
        self.assertEqual(encabezados, [c['encabezado'] for c in self.modelo['columnas']])
        self.assertEqual(encabezados[:5], ['Pozo', 'Departamento', 'Cuenca', 'Profundidad (m)', 'Tipo de pozo'])
        self.assertTrue(all(c.font.bold for c in ws[1]))

    def test_primera_fila_congelada_y_autofiltro(self):
        ws = self.wb['Pozos']
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(ws.auto_filter.ref, 'A1:K%d' % (len(self.modelo['filas']) + 1))
        self.assertGreater(ws.column_dimensions['A'].width, 5)

    def test_ninguna_celda_es_formula(self):
        for ws in self.wb.worksheets:
            for fila in ws.iter_rows():
                for c in fila:
                    self.assertNotEqual(c.data_type, 'f', '%s!%s es formula: %r' % (ws.title, c.coordinate, c.value))

    def test_ninguna_parte_del_zip_tiene_formulas(self):
        with zipfile.ZipFile(self.ruta) as z:
            self.assertIsNone(z.testzip())
            for nombre in z.namelist():
                self.assertNotRegex(z.read(nombre).decode('utf-8'), r'<f[ >/]', nombre)

    def test_textos_intactos_acentos_emoji_y_saltos_de_linea(self):
        ws = self.wb['Pozos']
        columnas = [c['clave'] for c in self.modelo['columnas']]
        idx_cuenca = columnas.index('cuenca')
        for r, fila in enumerate(self.modelo['filas'], start=2):
            leido = ws.cell(row=r, column=idx_cuenca + 1).value
            esperado = esperado_texto(fila['cuenca'])
            self.assertEqual(leido if leido is not None else '', esperado, 'fila %d' % r)
        todos = [ws.cell(row=r, column=idx_cuenca + 1).value for r in range(2, len(self.modelo['filas']) + 2)]
        self.assertIn('=1+1', todos)
        self.assertIn('+1+1', todos)
        self.assertIn('@SUM(1,2)', todos)
        self.assertIn('=2+2', todos)  # el modelo recorta espacios y tabulaciones de los extremos; sigue siendo texto
        self.assertIn('linea1\nlinea2', todos)
        self.assertIn('dos\nlineas', todos)
        self.assertIn('Ñandú, Río & "ñoño" <b> 🚰', todos)
        self.assertIn('ctrlx', todos)

    def test_textos_hostiles_son_texto_no_numero_ni_fecha(self):
        ws = self.wb['Pozos']
        columnas = [c['clave'] for c in self.modelo['columnas']]
        idx = columnas.index('cuenca') + 1
        for r in range(2, len(self.modelo['filas']) + 2):
            c = ws.cell(row=r, column=idx)
            if c.value is not None:
                self.assertEqual(c.data_type, 's', '%r' % c.value)
        valores = {ws.cell(row=r, column=idx).value: ws.cell(row=r, column=idx).data_type for r in range(2, len(self.modelo['filas']) + 2)}
        for texto in ('1/2', '1-2', '00123', '1e5', '31/12', 'TRUE', '#N/A'):
            self.assertEqual(valores[texto], 's', texto)

    def test_numeros_son_numericos_y_exactos(self):
        ws = self.wb['Pozos']
        columnas = [c['clave'] for c in self.modelo['columnas']]
        for clave in ('profundidad_m', 'latitud', 'longitud'):
            col = columnas.index(clave) + 1
            for r, fila in enumerate(self.modelo['filas'], start=2):
                c = ws.cell(row=r, column=col)
                esperado = fila[clave]
                if esperado is None or esperado == '':
                    self.assertIsNone(c.value, '%s fila %d' % (clave, r))
                else:
                    self.assertEqual(c.data_type, 'n', '%s fila %d' % (clave, r))
                    self.assertIsInstance(c.value, (int, float))
                    self.assertEqual(c.value, esperado)

    def test_coordenadas_negativas_con_formato_de_6_decimales(self):
        ws = self.wb['Pozos']
        columnas = [c['clave'] for c in self.modelo['columnas']]
        for clave in ('latitud', 'longitud'):
            col = columnas.index(clave) + 1
            c = ws.cell(row=2, column=col)
            self.assertLess(c.value, 0)
            self.assertEqual(c.number_format, '0.000000')

    def test_vacios_son_vacios_y_el_pozo_ausente_dice_sin_ubicacion_operativa(self):
        ws = self.wb['Pozos']
        columnas = [c['clave'] for c in self.modelo['columnas']]
        ultima = len(self.modelo['filas']) + 1
        self.assertEqual(ws.cell(row=ultima, column=1).value, '99-9999')
        self.assertEqual(ws.cell(row=ultima, column=columnas.index('estado_ubicacion') + 1).value, 'Sin ubicación operativa')
        for clave in ('cuenca', 'profundidad_m', 'tipo_pozo', 'latitud', 'longitud', 'tiene_itf'):
            self.assertIsNone(ws.cell(row=ultima, column=columnas.index(clave) + 1).value, clave)

    def test_itf_desconocido_queda_vacio_y_no_dice_no(self):
        ws = self.wb['Pozos']
        columnas = [c['clave'] for c in self.modelo['columnas']]
        col = columnas.index('tiene_itf') + 1
        valores = [ws.cell(row=r, column=col).value for r in range(2, len(self.modelo['filas']) + 2)]
        self.assertIn(None, valores)
        self.assertIn('Sí', valores)
        self.assertIn('No', valores)
        for r, fila in enumerate(self.modelo['filas'], start=2):
            self.assertEqual(ws.cell(row=r, column=col).value, fila['tiene_itf'] or None)

    def test_hoja_informacion(self):
        ws = self.wb['Información']
        filas = [[c.value for c in f] for f in ws.iter_rows()]
        self.assertEqual(filas[0], ['Dato', 'Valor'])
        self.assertEqual(filas[1:], [[k, v] for k, v in self.modelo['informacion']])
        info = dict(filas[1:])
        self.assertEqual(info['Generado (UTC)'], '2026-10-26T14:03:09Z')
        self.assertEqual(info['Generado (Mendoza)'], '2026-10-26 11:03:09 (Mendoza, UTC-3)')
        self.assertIsInstance(info['Generado (UTC)'], str)
        self.assertEqual(info['Cantidad de pozos'], len(self.modelo['filas']))
        self.assertEqual(info['Origen'], 'Mi selección')
        self.assertEqual(info['Búsqueda aplicada'], '=BUSQUEDA')              # texto literal, no formula
        self.assertEqual(info['Filtro de aptitud'], 'Apto')
        self.assertIn('Aptitud para reemplazo', info['Columnas incluidas'])
        for f in ws.iter_rows():
            for c in f:
                self.assertNotEqual(c.data_type, 'f')

    def test_nombre_de_archivo(self):
        self.assertEqual(self.modelo['nombreArchivo'], 'dgi-pozos-seleccion-2026-10-26.xlsx')

    def test_sin_datos_personales_en_ninguna_parte(self):
        with zipfile.ZipFile(self.ruta) as z:
            todo = ''.join(z.read(n).decode('utf-8') for n in z.namelist()).lower()
        for prohibido in ('titular', 'nc16', 'sessiontoken', 'clientrequestid', 'drivefileid'):
            self.assertNotIn(prohibido, todo)


@unittest.skipIf(openpyxl is None or NODE is None, 'faltan openpyxl o node')
class TestXlsxExportacionEscala(unittest.TestCase):
    def test_13804_pozos_se_abre_completo(self):
        with tempfile.TemporaryDirectory(prefix='xlsx_exp_') as d:
            ruta, modelo = generar(d, '13804')
            wb = openpyxl.load_workbook(ruta, read_only=True)
            ws = wb['Pozos']
            filas = list(ws.iter_rows(values_only=True))
            self.assertEqual(len(filas), 13804 + 1)
            self.assertEqual(filas[0][0], 'Pozo')
            self.assertEqual(len(filas[1]), 11)
            self.assertTrue(all(isinstance(f[3], (int, float)) or f[3] is None for f in filas[1:]))
            self.assertEqual(wb['Información']['B4'].value, 13804)
            wb.close()


if __name__ == '__main__':
    unittest.main()
